import streamlit as st
import pdfplumber
import anthropic
import os
import io
from dotenv import load_dotenv

load_dotenv()

# ══════════════════════════════════════════
#  إعداد الصفحة
# ══════════════════════════════════════════
st.set_page_config(
    page_title="محلل العقود الذكي",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ══════════════════════════════════════════
#  CSS — ديزاين مودرن داكن
# ══════════════════════════════════════════
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@300;400;500;700;800&display=swap');

/* ── Base ─────────────────────────────── */
html, body, [class*="css"] {
    font-family: 'Tajawal', sans-serif !important;
    direction: rtl;
}
.stApp { background: #080c14; }
section[data-testid="stSidebar"] { background: #0d1117 !important; border-left: 1px solid #1e2d42; }

/* ── Hero ─────────────────────────────── */
.hero {
    background: linear-gradient(135deg, #0d1b2e 0%, #0a1628 50%, #0d1b2e 100%);
    border: 1px solid #1a3a5c;
    border-radius: 20px;
    padding: 3rem 2rem;
    text-align: center;
    margin-bottom: 2rem;
    position: relative;
    overflow: hidden;
}
.hero::before {
    content: '';
    position: absolute;
    top: -60px; left: 50%;
    transform: translateX(-50%);
    width: 300px; height: 300px;
    background: radial-gradient(circle, rgba(56,139,253,0.12) 0%, transparent 70%);
    pointer-events: none;
}
.hero-icon { font-size: 3.5rem; margin-bottom: 1rem; display: block; }
.hero h1 {
    font-size: 2.4rem;
    font-weight: 800;
    color: #e6edf3;
    margin: 0 0 0.5rem;
    letter-spacing: -0.5px;
}
.hero p { color: #7d8590; font-size: 1.05rem; margin: 0; }
.hero-badge {
    display: inline-block;
    background: rgba(56,139,253,0.15);
    border: 1px solid rgba(56,139,253,0.3);
    color: #58a6ff;
    padding: 4px 14px;
    border-radius: 20px;
    font-size: 12px;
    font-weight: 600;
    margin-bottom: 1.2rem;
    letter-spacing: 0.5px;
}

/* ── Upload Zone ──────────────────────── */
div[data-testid="stFileUploader"] > div {
    background: #0d1117 !important;
    border: 2px dashed #1e3a5c !important;
    border-radius: 16px !important;
    padding: 2rem !important;
    transition: border-color 0.2s;
}
div[data-testid="stFileUploader"] > div:hover { border-color: #388bfd !important; }
div[data-testid="stFileUploader"] label { color: #8b949e !important; font-size: 14px !important; }

/* ── Stat Cards ───────────────────────── */
.stat-row { display: flex; gap: 12px; margin: 1.2rem 0; }
.stat-card {
    flex: 1;
    background: #0d1117;
    border: 1px solid #1e2d42;
    border-radius: 12px;
    padding: 1rem 1.2rem;
    text-align: center;
}
.stat-card .val { font-size: 1.7rem; font-weight: 700; color: #58a6ff; line-height: 1; }
.stat-card .lbl { font-size: 12px; color: #7d8590; margin-top: 4px; }

/* ── File Type Badge ──────────────────── */
.file-type-badge {
    display: inline-block;
    padding: 4px 12px;
    border-radius: 8px;
    font-size: 12px;
    font-weight: 700;
    letter-spacing: 0.5px;
    margin-bottom: 1rem;
}
.badge-pdf { background: #3d1a1a; color: #f87171; border: 1px solid #7f1d1d; }
.badge-docx { background: #1a2d4a; color: #60a5fa; border: 1px solid #1e3a6e; }
.badge-txt  { background: #1a2e1a; color: #4ade80; border: 1px solid #14532d; }
.badge-xlsx { background: #1e2d1a; color: #86efac; border: 1px solid #166534; }
.badge-img  { background: #2d1a3d; color: #c084fc; border: 1px solid #581c87; }

/* ── Analysis Select Tabs ─────────────── */
.type-grid { display: flex; gap: 10px; margin: 1rem 0 1.5rem; }
.type-card {
    flex: 1;
    background: #0d1117;
    border: 1px solid #1e2d42;
    border-radius: 12px;
    padding: 0.9rem;
    text-align: center;
    cursor: pointer;
    transition: all 0.2s;
}
.type-card:hover { border-color: #388bfd; background: #0d1b2e; }
.type-card.active { border-color: #388bfd; background: rgba(56,139,253,0.1); }
.type-card .ico { font-size: 1.4rem; }
.type-card .ttl { font-size: 12px; font-weight: 600; color: #e6edf3; margin-top: 4px; }

/* ── Result Box ───────────────────────── */
.result-wrap {
    background: #0d1117;
    border: 1px solid #1e2d42;
    border-radius: 16px;
    padding: 2rem;
    margin-top: 1.5rem;
    color: #c9d1d9;
    font-size: 15px;
    line-height: 1.85;
    white-space: pre-wrap;
}
.result-header {
    display: flex;
    align-items: center;
    gap: 10px;
    margin-bottom: 1.2rem;
    padding-bottom: 1rem;
    border-bottom: 1px solid #1e2d42;
}
.result-header span { font-size: 1rem; font-weight: 700; color: #e6edf3; }

/* ── Buttons ──────────────────────────── */
.stButton > button {
    background: linear-gradient(135deg, #1f6feb, #388bfd) !important;
    color: #fff !important;
    border: none !important;
    border-radius: 10px !important;
    font-size: 15px !important;
    font-weight: 700 !important;
    padding: 0.65rem 1.5rem !important;
    width: 100%;
    transition: opacity 0.2s !important;
    letter-spacing: 0.3px !important;
}
.stButton > button:hover { opacity: 0.88 !important; }

/* ── Sidebar ──────────────────────────── */
.sidebar-label {
    font-size: 11px;
    font-weight: 700;
    color: #7d8590;
    letter-spacing: 0.8px;
    text-transform: uppercase;
    margin: 1.2rem 0 0.4rem;
}
.sidebar-tip {
    background: #0d1b2e;
    border: 1px solid #1a3a5c;
    border-radius: 10px;
    padding: 0.8rem 1rem;
    font-size: 13px;
    color: #7d8590;
    line-height: 1.7;
    margin-top: 1rem;
}
.sidebar-tip strong { color: #58a6ff; }

/* ── Misc ─────────────────────────────── */
.stAlert { border-radius: 12px !important; }
.stSpinner > div { border-top-color: #388bfd !important; }
div[data-testid="stExpander"] {
    background: #0d1117;
    border: 1px solid #1e2d42;
    border-radius: 12px;
}
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════
#  دوال استخراج النص
# ══════════════════════════════════════════
def extract_from_pdf(f) -> str:
    with pdfplumber.open(f) as pdf:
        return "\n\n".join(p.extract_text() or "" for p in pdf.pages)

def extract_from_docx(f) -> str:
    from docx import Document
    doc = Document(f)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def extract_from_txt(f) -> str:
    raw = f.read()
    for enc in ("utf-8", "utf-16", "cp1256", "iso-8859-6"):
        try:
            return raw.decode(enc)
        except Exception:
            continue
    return raw.decode("utf-8", errors="replace")

def extract_from_xlsx(f) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(f, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[ورقة: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                lines.append(row_text)
    return "\n".join(lines)

def extract_from_image(f) -> str:
    """إرسال الصورة مباشرة لـ Claude للقراءة"""
    return "__IMAGE__" + f.read().hex() + "__EXT__" + f.name.split(".")[-1].lower()

def extract_text(uploaded_file) -> str:
    ext = uploaded_file.name.split(".")[-1].lower()
    uploaded_file.seek(0)
    if ext == "pdf":
        return extract_from_pdf(uploaded_file)
    elif ext in ("docx", "doc"):
        return extract_from_docx(uploaded_file)
    elif ext == "txt":
        return extract_from_txt(uploaded_file)
    elif ext in ("xlsx", "xls"):
        return extract_from_xlsx(uploaded_file)
    elif ext in ("png", "jpg", "jpeg", "webp"):
        return extract_from_image(uploaded_file)
    return ""

def get_file_badge(name: str) -> str:
    ext = name.split(".")[-1].lower()
    mapping = {
        "pdf": ("badge-pdf", "PDF"),
        "docx": ("badge-docx", "DOCX"), "doc": ("badge-docx", "DOC"),
        "txt": ("badge-txt", "TXT"),
        "xlsx": ("badge-xlsx", "XLSX"), "xls": ("badge-xlsx", "XLS"),
        "png": ("badge-img", "PNG"), "jpg": ("badge-img", "JPG"),
        "jpeg": ("badge-img", "JPEG"), "webp": ("badge-img", "WEBP"),
    }
    cls, label = mapping.get(ext, ("badge-txt", ext.upper()))
    return f'<span class="file-type-badge {cls}">{label}</span>'


# ══════════════════════════════════════════
#  دالة التحليل
# ══════════════════════════════════════════
def analyze_contract(text: str, analysis_type: str) -> str:
    api_key = os.getenv("ANTHROPIC_API_KEY") or st.session_state.get("api_key", "")
    if not api_key:
        return "❌ أدخل API Key أولاً من القائمة الجانبية"

    client = anthropic.Anthropic(api_key=api_key)

    prompts = {
        "تحليل شامل": """أنت محلل قانوني متخصص في العقود التجارية السعودية. حلّل الوثيقة التالية وأعطني:

## 1. 📋 ملخص الوثيقة
- نوع الوثيقة وموضوعها
- الأطراف المعنية
- المدة والقيمة المالية (إن وجدت)

## 2. ✅ أبرز الالتزامات
الطرف الأول والثاني بشكل واضح

## 3. ⚠️ بنود تستحق الانتباه
كل بند غير متوازن مع السبب

## 4. 🔴 مخاطر محتملة
المخاطر القانونية أو المالية

## 5. 💡 التوصية النهائية
هل الوثيقة متوازنة؟ ما أبرز نقاط التفاوض؟""",

        "البنود التعسفية": """أنت محلل قانوني. ركّز فقط على البنود التعسفية أو غير المتوازنة.
لكل بند:
- اقتبس البند بالنص
- اشرح لماذا هو تعسفي
- اقترح صياغة بديلة

إذا لم توجد بنود تعسفية، قل ذلك صراحةً.""",

        "الالتزامات المالية": """أنت محلل مالي. استخرج كل ما يتعلق بالمال:
- المبالغ المذكورة
- جدول الدفع
- الغرامات والجزاءات
- شروط الإلغاء والاسترداد
- أي التزامات مالية ضمنية

قدّم الأرقام بجداول إن أمكن."""
    }

    # معالجة الصور
    if text.startswith("__IMAGE__"):
        hex_data = text[9:text.index("__EXT__")]
        ext = text[text.index("__EXT__")+7:]
        img_bytes = bytes.fromhex(hex_data)
        import base64
        b64 = base64.b64encode(img_bytes).decode()
        media_map = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp"}
        media_type = media_map.get(ext, "image/jpeg")

        message = client.messages.create(
            model="claude-opus-4-6",
            max_tokens=2000,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                    {"type": "text", "text": prompts[analysis_type] + "\n\nحلّل الوثيقة في الصورة أعلاه."}
                ]
            }]
        )
    else:
        message = client.messages.create(
            model="claude-opus-4-6",
            max_tokens=2000,
            messages=[{
                "role": "user",
                "content": prompts[analysis_type] + f"\n\n---\nنص الوثيقة:\n{text[:8000]}"
            }]
        )

    return message.content[0].text


# ══════════════════════════════════════════
#  Sidebar
# ══════════════════════════════════════════
with st.sidebar:
    st.markdown('<div style="text-align:center;font-size:2rem;margin-bottom:0.5rem">⚖️</div>', unsafe_allow_html=True)
    st.markdown('<div style="text-align:center;font-weight:800;color:#e6edf3;font-size:1.1rem;margin-bottom:1.5rem">محلل العقود الذكي</div>', unsafe_allow_html=True)

    st.markdown('<div class="sidebar-label">🔑 مفتاح API</div>', unsafe_allow_html=True)
    api_key_input = st.text_input(
        "", type="password", placeholder="sk-ant-...",
        label_visibility="collapsed"
    )
    if api_key_input:
        st.session_state["api_key"] = api_key_input
        st.success("✅ تم حفظ المفتاح")

    st.markdown('<div class="sidebar-label">🔍 نوع التحليل</div>', unsafe_allow_html=True)
    analysis_type = st.selectbox(
        "", ["تحليل شامل", "البنود التعسفية", "الالتزامات المالية"],
        label_visibility="collapsed"
    )

    st.markdown("""<div class="sidebar-tip">
<strong>الصيغ المدعومة:</strong><br>
📄 PDF &nbsp;·&nbsp; 📝 DOCX &nbsp;·&nbsp; 📃 TXT<br>
📊 XLSX &nbsp;·&nbsp; 🖼️ PNG / JPG
<br><br>
<strong>للحصول على API Key:</strong><br>
console.anthropic.com
</div>""", unsafe_allow_html=True)


# ══════════════════════════════════════════
#  Hero
# ══════════════════════════════════════════
st.markdown("""
<div class="hero">
  <span class="hero-badge">مدعوم بـ Claude AI</span>
  <span class="hero-icon">⚖️</span>
  <h1>محلل العقود الذكي</h1>
  <p>ارفع أي وثيقة أو عقد واحصل على تحليل قانوني فوري وشامل</p>
</div>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════
#  رفع الملف
# ══════════════════════════════════════════
uploaded_file = st.file_uploader(
    "ارفع الوثيقة",
    type=["pdf", "docx", "doc", "txt", "xlsx", "xls", "png", "jpg", "jpeg", "webp"],
    help="يدعم PDF · Word · Excel · TXT · صور"
)

if uploaded_file:
    ext = uploaded_file.name.split(".")[-1].lower()
    badge_html = get_file_badge(uploaded_file.name)
    st.markdown(badge_html, unsafe_allow_html=True)

    with st.spinner("📖 جاري قراءة الملف..."):
        contract_text = extract_text(uploaded_file)

    is_image = contract_text.startswith("__IMAGE__")

    if not is_image:
        num_words = len(contract_text.split())
        num_chars = len(contract_text)
        st.markdown(f"""
        <div class="stat-row">
          <div class="stat-card"><div class="val">{uploaded_file.size // 1024} KB</div><div class="lbl">حجم الملف</div></div>
          <div class="stat-card"><div class="val">{num_words:,}</div><div class="lbl">كلمة</div></div>
          <div class="stat-card"><div class="val">{num_chars:,}</div><div class="lbl">حرف</div></div>
        </div>
        """, unsafe_allow_html=True)

        if num_chars < 50:
            st.error("⚠️ لم يُستخرج نص كافٍ. إذا كان الملف ممسوحاً ضوئياً، ارفعه كصورة PNG/JPG.")
            st.stop()

        with st.expander("👁️ معاينة النص المستخرج"):
            st.text(contract_text[:2000] + ("..." if len(contract_text) > 2000 else ""))
    else:
        st.info("🖼️ تم تحميل الصورة — سيقرأها الذكاء الاصطناعي مباشرة")

    st.markdown("<br>", unsafe_allow_html=True)

    if st.button(f"🔍 تحليل الآن — {analysis_type}", type="primary"):
        if not (os.getenv("ANTHROPIC_API_KEY") or st.session_state.get("api_key")):
            st.error("❌ أدخل API Key في القائمة الجانبية")
        else:
            with st.spinner("🤖 جاري التحليل بالذكاء الاصطناعي..."):
                result = analyze_contract(contract_text, analysis_type)

            st.markdown(f"""
            <div class="result-wrap">
              <div class="result-header">
                <span>📋 نتائج التحليل — {analysis_type}</span>
              </div>
              {result.replace(chr(10), '<br>')}
            </div>
            """, unsafe_allow_html=True)

            st.download_button(
                label="⬇️ تحميل التقرير",
                data=result,
                file_name=f"تحليل_{uploaded_file.name}.txt",
                mime="text/plain",
                use_container_width=True
            )
else:
    st.markdown("""
    <div style="text-align:center; padding: 3rem; color: #3d444d;">
      <div style="font-size:3rem;margin-bottom:1rem">📂</div>
      <div style="font-size:1rem">ارفع ملفك للبدء</div>
      <div style="font-size:13px;margin-top:0.5rem">PDF · DOCX · TXT · XLSX · صور</div>
    </div>
    """, unsafe_allow_html=True)